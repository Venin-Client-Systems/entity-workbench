"""Developer-only fixed synthetic DOCX render evidence; never a product dependency."""
from __future__ import annotations

import argparse
import hashlib
import json
import platform
import subprocess
import sys
import uuid
import zipfile
from datetime import datetime, timezone
from pathlib import Path
from xml.etree import ElementTree as ET

ROOT = Path(__file__).resolve().parents[1]
SOURCES = [
    "Cargo.lock", "crates/core/Cargo.toml", "crates/core/src/lib.rs",
    "crates/core/src/report_document.rs", "crates/core/src/report_docx.rs",
    "crates/core/src/report.rs", "crates/core/src/analytics.rs", "crates/core/src/domain.rs",
    "crates/core/examples/report_docx.rs", "crates/core/examples/support/report_fixture.rs",
    "crates/core/tests/report_docx.rs", "scripts/test_docx_report.py",
    "scripts/test_docx_evidence.py",
]
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
PARTS = [
    "[Content_Types].xml", "_rels/.rels", "docProps/core.xml", "word/document.xml",
    "word/styles.xml", "word/_rels/document.xml.rels",
]


def digest(path: Path) -> str:
    value = hashlib.sha256()
    with path.open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            value.update(chunk)
    return value.hexdigest()


def save(directory: Path, report: dict) -> None:
    # Keep every observation even if the host clock repeats or a retry fails.
    with (directory / f"observation-{uuid.uuid4().hex}.json").open("x") as target:
        json.dump(report, target, indent=2)
        target.write("\n")
    pending = directory / f"pending-{uuid.uuid4().hex}.json"
    pending.write_text(json.dumps(report, indent=2) + "\n")
    pending.replace(directory / "report.json")


def run_command(
    args: list[str], directory: Path, name: str, timeout: int = 180
) -> subprocess.CompletedProcess:
    with (directory / f"{name}.txt").open("w") as log:
        return subprocess.run(
            args, cwd=ROOT, stdout=log, stderr=subprocess.STDOUT,
            timeout=timeout, check=True,
        )


def inspect(docx: Path, frozen: Path) -> dict:
    if docx.stat().st_size > 32 * 1024 * 1024:
        raise ValueError("oversized DOCX")
    with zipfile.ZipFile(docx) as archive:
        if archive.namelist() != PARTS:
            raise ValueError("unexpected package entries or order")
        roots = {}
        for member in archive.infolist():
            if member.file_size > 32 * 1024 * 1024 or member.compress_type != zipfile.ZIP_STORED:
                raise ValueError("unexpected ZIP bounds or compression")
            raw = archive.read(member)
            if b"<!DOCTYPE" in raw or b"<!ENTITY" in raw:
                raise ValueError("unexpected XML declarations")
            roots[member.filename] = ET.fromstring(raw)
        for name in ["_rels/.rels", "word/_rels/document.xml.rels"]:
            for relation in roots[name]:
                if "TargetMode" in relation.attrib or relation.attrib["Target"] not in [
                    "word/document.xml", "docProps/core.xml", "styles.xml"
                ]:
                    raise ValueError("external or unknown relationship")
    main = roots["word/document.xml"]
    names = [node.attrib[W + "name"] for node in main.iter(W + "bookmarkStart")]
    links = [node.attrib[W + "anchor"] for node in main.iter(W + "hyperlink")]
    if len(names) != len(set(names)) or not set(links).issubset(names):
        raise ValueError("invalid citation bookmarks")
    text = "".join(
        (node.text or "") if node.tag == W + "t"
        else "\n" if node.tag == W + "br"
        else "\t" if node.tag == W + "tab"
        else "" for node in main.iter()
    )
    data = json.loads(frozen.read_text())
    for expected in [
        "<script>literal hostile text</script>", "000047",
        data["content"]["findings"][0]["assessment"], "copied-source-family-01",
    ]:
        if expected not in text:
            raise ValueError("missing retained specimen text")
    # Independent authoring-library read; it is a QA dependency only.
    from docx import Document
    opened = Document(docx)
    if len(opened.tables) != 1 or len(opened.tables[0].rows) != 19:
        raise ValueError("editable table structure changed")
    return {
        "parts": len(PARTS), "bookmarks": len(names), "links": len(links),
        "editable_table_rows": len(opened.tables[0].rows),
        "paragraphs": len(opened.paragraphs),
    }


def run(runtime: Path, output: Path) -> tuple[Path, bool]:
    run_id = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ") + "-" + uuid.uuid4().hex
    directory = output / run_id
    directory.mkdir(parents=True, exist_ok=False)
    report = {
        "schema_version": 1, "phase": "preflight", "outcome": "running",
        "visual_review": "not_performed", "native_word_review": "not_performed",
    }
    save(directory, report)
    try:
        python = runtime / "dependencies/python/bin/python3"
        if Path(sys.executable).resolve() != python.resolve():
            raise ValueError("use the selected bundled Python")
        renderer = runtime / "plugins/openai-primary-runtime/plugins/documents/skills/documents/render_docx.py"
        soffice = runtime / "dependencies/bin/override/soffice"
        native = runtime / "dependencies/native/libreoffice-headless/libreoffice/LibreOfficeDev.app/Contents/MacOS/soffice"
        report["runtime"] = {
            "manifest": json.loads((runtime / "runtime.json").read_text()),
            "manifest_sha256": digest(runtime / "runtime.json"),
            "renderer_sha256": digest(renderer),
            "soffice_wrapper_sha256": digest(soffice),
            "soffice_native_sha256": digest(native),
            "actual_version": subprocess.check_output(
                [str(soffice), "--version"], timeout=30, text=True
            ).strip(),
        }
        report["host"] = {
            "system": platform.system(), "machine": platform.machine(),
            "python": platform.python_version(),
        }
        report["source"] = {
            "commit": subprocess.check_output(
                ["git", "rev-parse", "HEAD"], cwd=ROOT, text=True
            ).strip(),
            "dirty": bool(subprocess.check_output(
                ["git", "status", "--porcelain"], cwd=ROOT, text=True
            ).strip()),
            "sha256": {name: digest(ROOT / name) for name in SOURCES},
        }
        report["phase"] = "generate"
        save(directory, report)
        for name in ["first", "repeat"]:
            run_command([
                "cargo", "run", "--locked", "-p", "workbench-core",
                "--example", "report_docx", "--", str(directory / name),
            ], directory, name, 300)
        first, repeat = directory / "first", directory / "repeat"
        for name in ["assessment.json", "assessment.docx"]:
            if digest(first / name) != digest(repeat / name):
                raise ValueError("non-deterministic frozen fixture or DOCX")
        report["phase"] = "structure"
        report["structure"] = inspect(first / "assessment.docx", first / "assessment.json")
        report["binary_sha256"] = digest(ROOT / "target/debug/examples/report_docx")
        report["artifacts"] = {
            name: {"sha256": digest(first / name), "bytes": (first / name).stat().st_size}
            for name in ["assessment.json", "assessment.docx"]
        }
        save(directory, report)
        report["phase"] = "render"
        save(directory, report)
        run_command([
            str(python), str(renderer), str(first / "assessment.docx"),
            "--output_dir", str(directory / "pages"), "--emit_pdf",
        ], directory, "render")
        pages = sorted(
            (directory / "pages").glob("page-*.png"),
            key=lambda path: int(path.stem.split("-")[-1]),
        )
        if not pages or len(pages) > 40:
            raise ValueError("unexpected synthetic page count")
        report["pages"] = [
            {"name": page.name, "sha256": digest(page), "bytes": page.stat().st_size}
            for page in pages
        ]
        report["outcome"] = "rendered_awaiting_visual_inspection"
        report["phase"] = "complete"
    except Exception as error:
        # Logs remain local. Do not leak paths or environment details into a public observation.
        report["outcome"] = "failed"
        report["error_type"] = type(error).__name__
        if isinstance(error, subprocess.CalledProcessError):
            report["return_code"] = error.returncode
    save(directory, report)
    return directory, report["outcome"] != "failed"


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--runtime-root", type=Path, required=True)
    parser.add_argument("--output", type=Path, default=ROOT / "artifacts/docx/runs")
    args = parser.parse_args()
    directory, passed = run(args.runtime_root.resolve(), args.output.resolve())
    print(directory)
    raise SystemExit(0 if passed else 1)
